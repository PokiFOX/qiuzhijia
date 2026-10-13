import datetime
import json
import mysql.connector
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import SheetDimension

from tapah import const
from tapah import reserved

# 与远端 Excel 模板命名样式一致（L/M 应可点，P/Q 等不应呈链接样式）
STYLE_NORMAL = "常规"
STYLE_HYPERLINK = "超链接"

def reset_data_rows(wb, sheet_name, first_row=2):
	ws = wb[sheet_name]
	for mrange in list(ws.merged_cells.ranges):
		if mrange.min_row >= first_row:
			ws.unmerge_cells(str(mrange))
	max_row = ws.max_row
	if max_row < first_row:
		return ws
	rows_to_drop = max_row - first_row + 1
	if rows_to_drop > 50000:
		max_col = max(ws.max_column or 1, 1)
		header = [ws.cell(1, c).value for c in range(1, max_col + 1)]
		idx = wb.sheetnames.index(sheet_name)
		wb.remove(ws)
		ws = wb.create_sheet(sheet_name, idx)
		for c, val in enumerate(header, 1):
			ws.cell(1, c).value = val
		return ws
	ws.delete_rows(first_row, rows_to_drop)
	return ws

def finalize_sheet(ws, last_data_row):
	last_data_row = max(int(last_data_row), 1)
	max_col = max(ws.max_column or 1, 1)
	if ws.max_row > last_data_row:
		ws.delete_rows(last_data_row + 1, ws.max_row - last_data_row)
	for row in list(ws.row_dimensions):
		if row > last_data_row:
			del ws.row_dimensions[row]
	ref = f"A1:{get_column_letter(max_col)}{last_data_row}"
	ws.sheet_dimension = SheetDimension(ref=ref)

def _looks_like_url(value) -> bool:
	text = str(value or "").strip()
	return text.startswith("http://") or text.startswith("https://")

def setcell(ws, col, row, value, as_url=False):
	cell = ws[f"{col}{row}"]
	if value is None or str(value).strip() == "":
		cell.value = None
		cell.hyperlink = None
		cell.style = STYLE_NORMAL
		return
	text = value if isinstance(value, str) else str(value)
	cell.value = text
	if as_url and _looks_like_url(text):
		url = text.strip()
		cell.hyperlink = url
		cell.style = STYLE_HYPERLINK
	else:
		cell.hyperlink = None
		cell.style = STYLE_NORMAL

conn = mysql.connector.connect(
	host		= reserved.mysql_host,
	port		= reserved.mysql_port,
	user		= reserved.mysql_username,
	password	= reserved.mysql_password,
	database	= reserved.mysql_database,
	autocommit	= True,
	auth_plugin = 'caching_sha2_password',
)
cursor = conn.cursor()

wb = openpyxl.load_workbook('企业列表.xlsx')

cursor.execute("SELECT id, zone FROM qzj_zone ORDER BY id")
zones = cursor.fetchall()
zone_name = {row[0]: row[1] for row in zones}

cursor.execute("SELECT id, level FROM qzj_level ORDER BY id")
levels = cursor.fetchall()
level_name = {row[0]: row[1] for row in levels}

cursor.execute("SELECT id, sector FROM qzj_sector ORDER BY id")
sectors = cursor.fetchall()
sector_name = {row[0]: row[1] for row in sectors}

cursor.execute("SELECT id, field, mapping, type, star, content FROM qzj_field ORDER BY id")
fields = cursor.fetchall()
field_name = {row[0]: row[1] for row in fields}

ent_name = {}
cursor.execute(
	"SELECT id, zone, city, name, shortname, brief, upper, level, sector, tag, website1, website2, icon, images, enttype, financial, growth, mapping, englishname FROM qzj_enterprise ORDER BY id"
)
enterprises = cursor.fetchall()
for row in enterprises:
	ent_name[row[0]] = row[3]

article1 = {}
article2 = {}
cursor.execute("SELECT enterprise_id, `index`, article FROM qzj_enterprise_article ORDER BY enterprise_id, `index`")
for ent_id, index, article in cursor.fetchall():
	if index == 1:
		article1.setdefault(ent_id, []).append(article)
	if index == 2:
		article2.setdefault(ent_id, []).append(article)

ws = reset_data_rows(wb, '全局设置')
n = max(len(zones), len(levels), len(sectors))
for i in range(n):
	row = i + 2
	if i < len(zones):
		ws.cell(row, 1).value = zones[i][1]
	if i < len(levels):
		ws.cell(row, 2).value = levels[i][1]
	if i < len(sectors):
		ws.cell(row, 3).value = sectors[i][1]
finalize_sheet(ws, 1 if n == 0 else n + 1)

ws = reset_data_rows(wb, '学科列表')
row = 2
for fid, name, mapping, type, star, content in fields:
	for mapname in (mapping or '').split(','):
		mapname = mapname.strip()
		if mapname == "": continue
		ws.cell(row, 1).value = name
		ws.cell(row, 2).value = mapname
		ws.cell(row, 3).value = type
		ws.cell(row, 4).value = star
		ws.cell(row, 5).value = content
		row += 1
finalize_sheet(ws, max(1, row - 1))

ws = reset_data_rows(wb, '第一批企业')
row = 2
for ent in enterprises:
	eid, zone_id, city, name, shortname, brief, upper, level_id, sector_id, tag, website1, website2, icon, images, enttype, financial, growth, mapping, englishname = ent
	setcell(ws, const.column_name, row, name)
	setcell(ws, const.column_zone, row, zone_name.get(zone_id, ''))
	setcell(ws, const.column_city, row, city)
	setcell(ws, const.column_short, row, shortname)
	setcell(ws, const.column_english, row, englishname or '')
	setcell(ws, const.column_brief, row, brief)
	setcell(ws, const.column_upper, row, upper)
	setcell(ws, const.column_sector, row, sector_name.get(sector_id, ''))
	setcell(ws, const.column_level, row, level_name.get(level_id, ''))
	setcell(ws, const.column_offer, row, mapping or '')
	setcell(ws, const.column_tag, row, (tag or '').strip())
	setcell(ws, const.column_website1, row, website1, as_url=True)
	setcell(ws, const.column_website2, row, website2, as_url=True)
	setcell(ws, const.column_icon, row, icon)
	setcell(ws, const.column_images, row, images)
	setcell(ws, const.column_enttype, row, enttype)
	setcell(ws, const.column_financial, row, financial)
	setcell(ws, const.column_growth, row, growth or '否')
	art1 = ','.join(article1.get(eid, []))
	art2 = ','.join(article2.get(eid, []))
	setcell(ws, const.column_article1, row, art1, as_url=_looks_like_url(art1.split(',')[0] if art1 else ''))
	setcell(ws, const.column_article2, row, art2, as_url=_looks_like_url(art2.split(',')[0] if art2 else ''))
	row += 1
finalize_sheet(ws, max(1, row - 1))

def stag_to_tag(stag):
	if stag == 1: return 'C9'
	if stag == 2: return '985'
	if stag == 3: return '211'
	if stag == 4: return '双非'
	if stag == 5: return '海外Top10'
	if stag == 6: return '海外Top50'
	if stag == 7: return '海外Top100'
	if stag == 8: return '其他海外院校'
	return ''

def detail_to_pipe(text):
	text = str(text or '').strip()
	if text == "": return ''
	if not text.startswith('['):
		return text
	items = json.loads(text)
	if not isinstance(items, list):
		return text
	parts = []
	for it in items:
		if not isinstance(it, dict): continue
		parts.append(f"{it.get('enterpriseName', '')}|{it.get('department', '')}|{it.get('position', '')}")
	return ','.join(parts)

cursor.execute(
	"SELECT id, name, enterprise, field, tags, student, school1, stag1, field1, school2, stag2, field2, year, detail, dep FROM qzj_case ORDER BY id"
)
cases = cursor.fetchall()

ws = reset_data_rows(wb, '成功案例')
row = 2
for case in cases:
	cid, name, ent_id, field_id, tags, student, school1, stag1, field1, school2, stag2, field2, year, detail, dep = case
	ws.cell(row, 1).value = name
	ws.cell(row, 2).value = ent_name.get(ent_id, '')
	ws.cell(row, 3).value = field_name.get(field_id, '')
	ws.cell(row, 4).value = tags
	ws.cell(row, 5).value = student
	ws.cell(row, 6).value = school1
	ws.cell(row, 7).value = stag_to_tag(stag1)
	ws.cell(row, 8).value = field1
	ws.cell(row, 9).value = school2
	ws.cell(row, 10).value = stag_to_tag(stag2)
	ws.cell(row, 11).value = field2
	ws.cell(row, 12).value = year
	ws.cell(row, 13).value = detail_to_pipe(detail)
	ws.cell(row, 14).value = dep
	row += 1
finalize_sheet(ws, max(1, row - 1))

cursor.close()
conn.close()

stamp = datetime.datetime.now().strftime('%Y%m%d%H%M%S')
out_name = f'企业列表_{stamp}.xlsx'
wb.save(out_name)
wb.close()
print(out_name)
